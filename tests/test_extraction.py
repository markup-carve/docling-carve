import pytest
from docling_carve import _native


def test_real_html_extraction_without_ocr_or_model_downloads(tmp_path):
    pytest.importorskip("docling")
    from docling_carve.extraction import convert_document

    source = tmp_path / "report.html"
    source.write_text(
        '<html><body><h1>Extracted title</h1><p>Extracted passage.</p><table><tr><th colspan="2">Limits</th></tr><tr><td>A</td><td>B</td></tr></table></body></html>'
    )
    result = convert_document(source, ocr=False)
    html = _native.to_html(result.value)
    assert "Extracted title" in html and "Extracted passage" in html
    assert 'colspan="2"' in html
    assert result.source["extraction_status"] == "success"
    assert result.document["tables"]
    with pytest.raises(ValueError, match="allow_url"):
        convert_document("https://example.com/report.pdf")
    with pytest.raises(ValueError, match="max_input_bytes"):
        convert_document(source, max_input_bytes=10)


def test_real_docx_extraction(tmp_path):
    pytest.importorskip("docling")
    from docx import Document
    from docling_carve.extraction import convert_document

    source = tmp_path / "office.docx"
    office = Document()
    office.add_heading("Office heading", 1)
    office.add_paragraph("Office passage.")
    office.save(source)
    result = convert_document(source, ocr=False)
    assert "Office heading" in result.value and "Office passage" in result.value
    assert result.source["extraction_status"] == "success"


def test_real_native_pdf_extraction(tmp_path):
    pytest.importorskip("docling")
    from docling_carve.extraction import convert_document

    # A one-page text PDF avoids OCR and layout model downloads in this test.
    stream = b"BT /F1 18 Tf 72 720 Td (Native PDF passage) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(str(index).encode() + b" 0 obj\n" + obj + b"\nendobj\n")
    start = len(pdf)
    pdf.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode())
    source = tmp_path / "native.pdf"
    source.write_bytes(pdf)
    result = convert_document(source, pdf_pipeline="native", ocr=False)
    assert "Native PDF passage" in result.value
    assert result.provenance[0]["pages"]
